"""
De behandelingen uit het teeltlogboek "Klimaatregistratie tuin N" (blad Overview) in de database: de historie
vóór Agro4all. Rekenwerk in logic/klimaatregistratie.py.

Twee stappen, omdat het Excel-bestand op OneDrive staat en de uurtaak op de NAS draait:
1. Op de pc het Excel-bestand omzetten naar een klein JSON-bestand in de map agro4all op de NAS:
       py klimaatregistratie_import.py --xlsx "C:\\...\\Klimaatregistratie tuin 3.xlsx" --tuin 3 \\
           --json \\\\192.168.0.197\\Appdata\\VEMteelt\\agro4all\\klimaatregistratie-tuin3.json
2. Elk uur (priva_bestanden.py) leest dit script alle klimaatregistratie-tuin*.json in die map en vervangt de
   behandelingen van die bron (klimaatregistratie-tuinN). Alleen dagen vóór de eerste Agro4all-behandeling in
   die tuin, zodat niets dubbel staat. De codes worden bij het inlezen vertaald (logic/klimaatregistratie.CODES),
   dus een verbeterde code werkt zonder stap 1 opnieuw te doen.

Gebruik (stap 2): python klimaatregistratie_import.py [--map /agro4all] [--droog] [--doel productie]
"""
import argparse
import json
import os
import sys
from datetime import date, datetime, timezone
from pathlib import Path


def lees_argumenten():
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("--xlsx", help="stap 1: het Excel-bestand (blad Overview)")
    parser.add_argument("--tuin", type=int, help="stap 1: het tuinnummer van dat bestand")
    parser.add_argument("--json", help="stap 1: waar het JSON-bestand komt")
    parser.add_argument("--map", default=os.environ.get("AGRO4ALL_MAP", "/agro4all"),
                        help="stap 2: map met klimaatregistratie-tuin*.json")
    parser.add_argument("--doel", choices=("app", "productie"), default="app",
                        help="app = DATABASE_URL (standaard), productie = DATABASE_URL_PRODUCTIE")
    parser.add_argument("--droog", action="store_true", help="alleen tonen, niets opslaan")
    return parser.parse_args()


def omzetten(args):
    """Stap 1: Excel → JSON (vak, datum, regel, codes per behandeling)."""
    from openpyxl import load_workbook
    from logic import klimaatregistratie as kr
    wb = load_workbook(args.xlsx, read_only=True, data_only=True)
    beh = kr.behandelingen(wb["Overview"].iter_rows(values_only=True))
    uit = {"tuin": args.tuin, "bron": f"klimaatregistratie-tuin{args.tuin}", "bestand": Path(args.xlsx).name,
           "gemaakt": datetime.now(timezone.utc).isoformat(timespec="seconds"),
           "behandelingen": [[v, d.isoformat(), r, c] for v, d, r, c in beh]}
    if not args.droog:
        pad = Path(args.json)
        tijdelijk = pad.with_suffix(".tmp")
        tijdelijk.write_text(json.dumps(uit, ensure_ascii=False), encoding="utf-8")
        tijdelijk.replace(pad)
    dagen = sorted(d for _, d, _, _ in beh)
    print(f"{len(beh)} behandelingen uit {Path(args.xlsx).name}"
          + (f" ({dagen[0]:%d-%m-%y} t/m {dagen[-1]:%d-%m-%y})" if dagen else "")
          + (" (droog)" if args.droog else f" → {args.json}"))


def inlezen(args):
    """Stap 2: elk JSON-bestand in de map naar de database."""
    stempel = f"{datetime.now(timezone.utc):%d-%m-%y %H:%M} UTC"
    bestanden = sorted(Path(args.map).glob("klimaatregistratie-tuin*.json"))
    if not bestanden:
        print(f"[{stempel}] Geen klimaatregistratie-tuin*.json in {args.map}; niets gedaan.")
        return
    if args.doel == "productie":
        doel_url = os.environ.get("DATABASE_URL_PRODUCTIE")
        if not doel_url:
            print("FOUT: DATABASE_URL_PRODUCTIE is niet gezet.", file=sys.stderr)
            sys.exit(1)
        os.environ["DATABASE_URL"] = doel_url

    import database
    from logic import agro4all as a4
    from logic import klimaatregistratie as kr

    tuinen = {t["nummer"]: t for t in database.get_tuinen()}
    if not args.droog:
        database.init_db()
    for pad in bestanden:
        data = json.loads(pad.read_text(encoding="utf-8"))
        tuin = tuinen[data["tuin"]]
        beh = [(v, date.fromisoformat(d), r, c) for v, d, r, c in data["behandelingen"]]
        # Vanaf de eerste Agro4all-behandeling in deze tuin geldt Agro4all.
        with database.get_connection() as conn:
            c = conn.cursor()
            c.execute("SELECT MIN(datum) FROM behandeling WHERE bron = %s AND tuin_id = %s", (a4.BRON, tuin["id"]))
            eerste = c.fetchone()[0]
        tot = date.fromisoformat(eerste) if eerste else None
        beh = [b for b in beh if tot is None or b[1] < tot]
        middelen, regels, onbekend = kr.regels(beh, data["tuin"], data["bron"])
        for r in regels:
            r["tuin_id"] = tuin["id"]
            r.pop("tuin_nummer")
        dagen = sorted({r["datum"] for r in regels})
        samenvatting = (f"{tuin['naam']}: {len(beh)} behandelingen, {len(regels)} regels, {len(middelen)} middelen"
                        + (f" ({dagen[0]:%d-%m-%y} t/m {dagen[-1]:%d-%m-%y})" if dagen else "")
                        + (f"; tot de eerste Agro4all-dag {tot:%d-%m-%y}" if tot else "")
                        + (f"; onbekende codes: {', '.join(f'{k} ({n}x)' for k, n in sorted(onbekend.items()))}"
                           if onbekend else ""))
        if args.droog:
            print(f"[{stempel}] DROOG, niets opgeslagen; zou opslaan: {samenvatting}")
            continue
        database.vervang_behandelingen(data["bron"], middelen, regels)
        print(f"[{stempel}] {samenvatting}")


def main():
    args = lees_argumenten()
    if args.xlsx:
        if not (args.tuin and args.json):
            print("FOUT: bij --xlsx horen ook --tuin en --json.", file=sys.stderr)
            sys.exit(1)
        omzetten(args)
    else:
        inlezen(args)


if __name__ == "__main__":
    main()
