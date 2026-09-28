"""
Historische teelten met hun klimaat, uit de Excel-kopieën en de database (alleen lezen).

Bronnen
- Teelt\\Klimaatregistratie tuin 1 kopie.xlsx en tuin 3 kopie.xlsx (tuin 3 alleen code S):
  blad Teelt (weekwaarden °C en lichtsom binnen per teelt), blad Stek (plantdatum),
  blad Aantekeningen (tuin 1: Florgib- en oogstdagen).
- De database (DATABASE_URL): teelten, teeltvakken, oogstregistraties, klimaatdata_dag.

Keuzes (afgesproken)
- Excel gaat voor waar Excel en database verschillen (licht afdeling 1 tuin 3).
- Excel kent alleen weekwaarden: elke dag van een week krijgt de weekwaarde.
- Oogstdatum: uit de database; bij tuin 3 vóór de app is die op de week nauwkeurig
  (±3 d). Een oogstdatum die precies de teeltduur-tabel is, is bij de import
  geschat en telt niet (dan de oogstweek uit Excel, of de teelt valt af).
"""
import os
import re
import sys
from collections import defaultdict
from datetime import date, timedelta

import openpyxl
import pandas as pd
import numpy as np

HIER = os.path.dirname(os.path.abspath(__file__))
PROJECT = os.path.dirname(HIER)
sys.path.insert(0, PROJECT)

EXCEL = {1: os.path.join(PROJECT, "Teelt", "Klimaatregistratie tuin 1 kopie.xlsx"),
         3: os.path.join(PROJECT, "Teelt", "Klimaatregistratie tuin 3 kopie.xlsx")}
TUIN3_CODE = "S"
WEEKKOP = re.compile(r"^(\d{4})-(\d{1,2})$")
DAGEN = {"ma": 1, "di": 2, "wo": 3, "do": 4, "vr": 5, "za": 6, "zo": 7}
NA_OOGST_DAGEN = 45      # klimaat na de oogst, voor als een model later uitkomt
MIN_DEKKING = 0.9        # deel van de teeltdagen met klimaat, anders valt de teelt af
DUUR_GRENZEN = (28, 140) # onwaarschijnlijke teeltduur valt af (datafout)


def _open(pad):
    """Werkmap openen; staat hij open in Excel (vergrendeld), dan via een tijdelijke kopie."""
    try:
        return openpyxl.load_workbook(pad, read_only=True, data_only=True)
    except PermissionError:
        import shutil
        import tempfile
        kopie = os.path.join(tempfile.gettempdir(), "vem_" + os.path.basename(pad))
        try:
            shutil.copyfile(pad, kopie)
        except PermissionError:
            # Excel vergrendelt soms ook lezen; Windows' eigen kopie lukt dan nog wel.
            import subprocess
            subprocess.run(["powershell", "-NoProfile", "-Command",
                            f"Copy-Item -LiteralPath '{pad}' -Destination '{kopie}' -Force"], check=True)
        return openpyxl.load_workbook(kopie, read_only=True, data_only=True)


def lees_excel_teelt(wb, tuin):
    """Blad Teelt: [{jaar, plantweek, vak, code, T: {(j, w): °C}, L: {(j, w): J/cm²}}]."""
    rows = list(wb["Teelt"].iter_rows(values_only=True))
    kop = rows[0]
    weekkol = {i: tuple(map(int, WEEKKOP.match(str(k)).groups())) for i, k in enumerate(kop)
               if k and WEEKKOP.match(str(k))}
    met_code = tuin == 3
    eenheid_i = 4 if met_code else 3
    teelten = defaultdict(dict)
    for r in rows[2:]:
        if r[0] is None or r[1] is None or r[2] is None:
            continue
        code = r[3] if met_code else None
        if met_code and code != TUIN3_CODE:
            continue
        reeks = {weekkol[i]: float(v) for i, v in enumerate(r)
                 if i in weekkol and isinstance(v, (int, float)) and v > 0}
        teelten[(int(r[1]), int(r[0]), int(r[2]), code)]["T" if "C" in str(r[eenheid_i]) else "L"] = reeks
    return [{"tuin": tuin, "jaar": j, "plantweek": w, "vak": v, "code": c, "T": d.get("T", {}), "L": d.get("L", {})}
            for (j, w, v, c), d in teelten.items() if d.get("T") or d.get("L")]


def lees_stek_plantdata(wb):
    """Blad Stek: {(jaar, week, vak): plantdatum}."""
    uit = {}
    if "Stek" not in wb.sheetnames:
        return uit
    rijen = wb["Stek"].iter_rows(values_only=True)
    next(rijen)
    for r in rijen:
        jaar, week, dag, vak = (list(r) + [None] * 4)[:4]
        if None in (jaar, week, dag, vak):
            continue
        dagnr = DAGEN.get(str(dag).strip().lower()[:2])
        try:
            uit[(int(jaar), int(week), int(vak))] = date.fromisocalendar(int(jaar), int(week), dagnr)
        except (TypeError, ValueError):
            continue
    return uit


def _dagweek(code, jaar, plantweek):
    """'wo-14' → datum (in het plantjaar, of het jaar erna als de week vóór de plantweek ligt)."""
    if not isinstance(code, str):
        return None
    m = re.match(r"\s*([a-zA-Z]{2})\w*\s*-\s*(\d{1,2})", code)
    if not m or m.group(1).lower() not in DAGEN:
        return None
    week = int(m.group(2))
    try:
        return date.fromisocalendar(jaar + (week < plantweek), week, DAGEN[m.group(1).lower()])
    except ValueError:
        return None


def lees_florgib_tuin1(wb):
    """Blad Aantekeningen (tuin 1): {(jaar, plantweek, vak): Florgib-datum}."""
    uit = {}
    if "Aantekeningen" not in wb.sheetnames:
        return uit
    rows = list(wb["Aantekeningen"].iter_rows(values_only=True))
    kop = [str(k).strip() if k else "" for k in rows[0]]
    for r in rows[1:]:
        if not isinstance(r[0], (int, float)) or not isinstance(r[1], (int, float)):
            continue
        d = dict(zip(kop, r))
        jaar, week = int(r[0]), int(r[1])
        florgib = _dagweek(d.get("Florgib"), jaar, week)
        vakken = [int(v) for v in re.findall(r"\d+", str(d.get("Tralie") or ""))]
        if "-" in str(d.get("Tralie")) and len(vakken) == 2:
            vakken = list(range(vakken[0], vakken[1] + 1))
        for vak in vakken:
            if florgib:
                uit[(jaar, week, vak)] = florgib
    return uit


def lees_florgib_overview(wb):
    """
    Blad Overview: {vak: [datums]} waarop in de GBM-rij "Fg" (Florgib) staat.
    De kop heeft per week een blok van zeven dagen dat op zondag begint (jaar,
    week en dagletter in de eerste drie rijen). Staat het jaartal niet goed,
    dan telt het jaar op zodra het weeknummer terugspringt.
    """
    if "Overview" not in wb.sheetnames:
        return {}
    rijen = wb["Overview"].iter_rows(values_only=True)
    jaren, weken, dagen = next(rijen), next(rijen), next(rijen)
    next(rijen)
    kolomdatum, jaar, vorige_week, positie, blokstart = {}, None, None, 0, None
    for i in range(len(dagen)):
        if i < len(weken) and isinstance(weken[i], (int, float)) and i < len(dagen) and dagen[i] in ("Z", "zo"):
            week = int(weken[i])
            kop_jaar = jaren[i] if i < len(jaren) and isinstance(jaren[i], (int, float)) else None
            if jaar is None:
                jaar = int(kop_jaar) if kop_jaar else None
            elif vorige_week is not None and week < vorige_week:
                jaar += 1
            elif kop_jaar and int(kop_jaar) > jaar:
                jaar = int(kop_jaar)
            vorige_week, positie = week, 0
            try:
                blokstart = date.fromisocalendar(jaar, week, 1) - timedelta(days=1) if jaar else None
            except ValueError:
                blokstart = None
        elif blokstart is not None:
            positie += 1
        if blokstart is not None and dagen[i]:
            kolomdatum[i] = blokstart + timedelta(days=positie)
    uit, vak = defaultdict(list), None
    for rij in rijen:
        if rij[0] is not None and isinstance(rij[0], (int, float)):
            vak = int(rij[0])
        if rij[1] != "GBM" or vak is None:
            continue
        for i, d in kolomdatum.items():
            waarde = rij[i] if i < len(rij) else None
            if isinstance(waarde, str) and re.search(r"\bfg\b", waarde.strip().lower()):
                uit[vak].append(d)
    return {v: sorted(ds) for v, ds in uit.items()}


def lees_database():
    """Teelten (met aantal oogstregistraties), vak→afdeling en klimaat per dag."""
    from database import get_connection, meting
    with get_connection() as c:
        cur = c.cursor()
        cur.execute("""
            SELECT t.id, tu.nummer, v.vaknummer, v.afdeling, t.datum_teelt_start, t.datum_half, t.lengte_half,
                   t.datum_oogst, t.lengte_eind, t.oogstgewicht,
                   (SELECT COUNT(*) FROM oogstregistraties o WHERE o.teelt_id = t.id)
            FROM teelten t JOIN teeltvakken v ON v.id = t.teeltvak_id JOIN tuinen tu ON tu.id = v.tuin_id
            WHERE v.vaknummer IS NOT NULL""")
        teelten = pd.DataFrame(cur.fetchall(), columns=[
            "id", "tuin", "vak", "afdeling", "start", "half", "lengte_half", "oogst", "lengte", "gewicht", "registraties"])
        cur.execute("""SELECT tu.nummer, v.vaknummer, v.afdeling FROM teeltvakken v JOIN tuinen tu ON tu.id = v.tuin_id
                       WHERE v.vaknummer IS NOT NULL""")
        vak_afd = {(t, v): a for t, v, a in cur.fetchall()}
        cur.execute("""SELECT tu.nummer, k.afdeling, k.datum, k.gem_temperatuur, k.stralingssom_dag
                       FROM klimaatdata_dag k JOIN tuinen tu ON tu.id = k.tuin_id""")
        klimaat = pd.DataFrame(cur.fetchall(), columns=["tuin", "afdeling", "datum", "T", "L"])
    for k in ("lengte", "gewicht", "lengte_half"):
        teelten[k] = teelten[k].map(meting)
    return teelten, vak_afd, klimaat


def _getal(waarde):
    """Getal, of None voor None/NaN (pandas maakt van None in een getalkolom NaN)."""
    if waarde is None:
        return None
    waarde = float(waarde)
    return None if waarde != waarde else waarde


def _dagen_van_week(jaar_week):
    maandag = date.fromisocalendar(jaar_week[0], jaar_week[1], 1)
    return [maandag + timedelta(days=i) for i in range(7)]


def bouw_dataset():
    """
    Alle bruikbare afgeronde teelten als dicts met: tuin, vak, afdeling, code,
    start, oogst, precisie ("dag"/"week"), florgib, gewicht, lengte, bron, duur
    (dagen), T/L (per dag van plantdag tot oogst) en T_ext/L_ext (tot
    NA_OOGST_DAGEN erna). Plus het afdelingsklimaat per dag en een logboek van
    wat er afviel.
    """
    from database import bereken_verwachte_oogstdatum

    db, vak_afd, db_klimaat = lees_database()
    excel, stek, florgib_t1, florgib_overview = [], {}, {}, {}
    for tuin, pad in EXCEL.items():
        wb = _open(pad)
        excel += lees_excel_teelt(wb, tuin)
        stek.update({(tuin,) + k: v for k, v in lees_stek_plantdata(wb).items()})
        florgib_overview.update({(tuin, vak): ds for vak, ds in lees_florgib_overview(wb).items()})
        if tuin == 1:
            florgib_t1 = lees_florgib_tuin1(wb)

    # Afdelingsklimaat per dag: database als basis, Excel (afdelingsgemiddelde per week) gaat voor.
    klimaat = defaultdict(dict)
    for r in db_klimaat.itertuples():
        if r.T is not None and r.L is not None:
            klimaat[(r.tuin, r.afdeling)][date.fromisoformat(r.datum)] = (float(r.T), float(r.L))
    afd_week = defaultdict(list)
    for t in excel:
        afd = vak_afd.get((t["tuin"], t["vak"]))
        for w in set(t["T"]) & set(t["L"]):
            afd_week[(t["tuin"], afd, w)].append((t["T"][w], t["L"][w]))
    for (tuin, afd, w), waarden in afd_week.items():
        T, L = np.mean([x[0] for x in waarden]), np.mean([x[1] for x in waarden])
        for d in _dagen_van_week(w):
            klimaat[(tuin, afd)][d] = (float(T), float(L))

    db["iso"] = db["start"].map(lambda s: date.fromisoformat(s).isocalendar()[:2])
    db_per_sleutel = {(r.tuin, r.iso[0], r.iso[1], r.vak): r for r in db.itertuples()}
    excel_per_sleutel = {(t["tuin"], t["jaar"], t["plantweek"], t["vak"]): t for t in excel}

    log = defaultdict(int)
    teelten = []
    for sleutel in sorted(set(excel_per_sleutel) | set(db_per_sleutel)):
        tuin, jaar, week, vak = sleutel
        ex, r = excel_per_sleutel.get(sleutel), db_per_sleutel.get(sleutel)
        if tuin == 3 and ex is None and r is not None and date.fromisoformat(r.start) < date(2025, 5, 12):
            continue  # tuin 3 vóór de app: alleen via Excel (code S)
        afd = vak_afd.get((tuin, vak))
        start = date.fromisoformat(r.start) if r is not None else stek.get(sleutel)
        start_precisie = "dag"
        if start is None:
            start, start_precisie = date.fromisocalendar(jaar, week, 3), "week"

        oogst, precisie = None, None
        if r is not None and isinstance(r.oogst, str):
            _, geschat = bereken_verwachte_oogstdatum(r.start)
            if geschat is not None and str(geschat) == r.oogst:
                log["oogstdatum in database is geschat (teeltduur-tabel)"] += 1
            else:
                oogst = date.fromisoformat(r.oogst)
                precisie = "dag" if (tuin == 1 or r.registraties) else "week"
        if oogst is None and ex is not None:
            weken = sorted(set(ex["T"]) | set(ex["L"]))
            laatste = weken[-1]
            if laatste < date.today().isocalendar()[:2]:
                oogst = date.fromisocalendar(laatste[0], laatste[1], start.isoweekday())
                precisie = "week"
        if oogst is None:
            log["nog lopend of geen oogst bekend"] += 1
            continue
        if start_precisie == "week" and precisie == "dag":
            precisie = "week"
        duur = (oogst - start).days
        if not DUUR_GRENZEN[0] <= duur <= DUUR_GRENZEN[1]:
            log[f"teeltduur buiten {DUUR_GRENZEN[0]}-{DUUR_GRENZEN[1]} dagen"] += 1
            continue

        # Dagreeks: eigen Excel-weken gaan voor, dan het afdelingsklimaat.
        eigen = {}
        if ex is not None:
            for w in set(ex["T"]) & set(ex["L"]):
                for d in _dagen_van_week(w):
                    eigen[d] = (ex["T"][w], ex["L"][w])
        dagen = [start + timedelta(days=i) for i in range(duur + NA_OOGST_DAGEN)]
        reeks = [eigen.get(d) or klimaat.get((tuin, afd), {}).get(d) for d in dagen]
        dekking = sum(1 for x in reeks[:duur] if x) / duur
        if dekking < MIN_DEKKING:
            log["te weinig klimaatdata (< 90 % van de teeltdagen)"] += 1
            continue
        # Kleine gaten opvullen met de vorige bekende dag; na de oogst mag de reeks eindigen.
        T, L, vorige = [], [], None
        for i, x in enumerate(reeks):
            if x is None and i >= duur:
                break
            x = x or vorige
            vorige = x
            T.append(x[0]); L.append(x[1])
        T, L = np.array(T), np.array(L)

        # Florgib: de registratie in de app, anders het eerste "Fg" in Overview
        # binnen de teelt, anders de Aantekeningen (tuin 1).
        fg_db = date.fromisoformat(r.half) if r is not None and isinstance(r.half, str) else None
        fg_ov = next((d for d in florgib_overview.get((tuin, vak), []) if start < d < oogst), None)
        fg = fg_db or fg_ov or (florgib_t1.get((jaar, week, vak)) if tuin == 1 else None)
        if fg_db and fg_ov:
            log["Florgib in app én Overview (verschil in dagen)"] = log.get(
                "Florgib in app én Overview (verschil in dagen)", []) + [(fg_db - fg_ov).days]
        if fg is not None and not (start < fg < oogst):
            log["Florgib-datum buiten de teelt (genegeerd)"] += 1
            fg = None
        bron_fg = None if fg is None else ("app" if fg == fg_db else "overview" if fg == fg_ov else "aantekeningen")

        # Weekwaarden voor de historietabel: Excel waar dat er is, anders het
        # weekgemiddelde uit de dagwaarden van de afdeling (database).
        weken = {}
        w = start - timedelta(days=start.weekday())
        while w <= oogst:
            jw = w.isocalendar()[:2]
            if ex is not None and jw in ex["T"] and jw in ex["L"]:
                weken[jw] = (ex["T"][jw], ex["L"][jw], "excel")
            else:
                dagwaarden = [klimaat.get((tuin, afd), {}).get(w + timedelta(days=i)) for i in range(7)]
                dagwaarden = [x for x in dagwaarden if x]
                if dagwaarden:
                    weken[jw] = (float(np.mean([x[0] for x in dagwaarden])), float(np.mean([x[1] for x in dagwaarden])),
                                 "database")
            w += timedelta(days=7)

        teelten.append({
            "tuin": tuin, "jaar": jaar, "plantweek": week, "vak": vak, "afdeling": afd,
            "start_precisie": start_precisie, "teelt_id": int(r.id) if r is not None else None, "weken": weken,
            "code": ex["code"] if ex else None, "start": start, "oogst": oogst, "precisie": precisie,
            "florgib": fg, "florgib_bron": bron_fg, "gewicht": _getal(r.gewicht) if r is not None else None,
            "lengte": _getal(r.lengte) if r is not None else None,
            "bron": "beide" if (ex is not None and r is not None) else ("excel" if ex is not None else "database"),
            "duur": float(duur), "T": T[:duur], "L": L[:duur], "T_ext": T, "L_ext": L,
        })
    return teelten, klimaat, dict(log)
