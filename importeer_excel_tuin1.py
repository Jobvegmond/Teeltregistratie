"""
Zet de teelthistorie van tuin 1 uit de oude Excel-registratie over naar de
database, zodat tuin 1 in de app niet leeg begint.

Bron: Teelt\\Klimaatregistratie tuin 1 kopie.xlsx
  - blad "Stek":          plantdatum (jaar/week/dag), vak, aantal planten,
                          ras en de stekbeoordeling
  - blad "Aantekeningen": per planting de vakken, de eerste en laatste
                          oogstdag en de lengte bij oogst
  - blad "Overview":      watergift per vak per dag (de rijen "W")

Eén planting, meerdere vakken
-----------------------------
Op tuin 1 wordt een planting in één week over twee tot vier vakken verdeeld.
De database kent een teelt per vak, dus zo'n planting wordt hier een teelt per
vak met dezelfde plantweek. Aan de code is te zien dat ze bij elkaar horen: die
begint met jaar + plantweek (2605..), alleen de laatste twee cijfers, het vak,
verschillen. De oogstdag en de lengte uit Aantekeningen gelden voor de hele
planting en komen dus bij elk vak van die planting te staan.

Wat wel en niet mee gaat
------------------------
  - alles vanaf 2025 week 49, de eerste planting van de huidige ronde;
  - teelten die nog lopen gaan gewoon mee, zonder oogstdatum, zodat de app de
    huidige stand van tuin 1 laat zien;
  - staat de laatste oogstdag er niet, dan komt de oogstdatum uit de eerste
    oogstdag plus de gebruikelijke oogstduur, en anders uit de oogstweek in
    blad Teelt (dat blad is t/m maart bijgehouden). Ligt de uitkomst in de
    toekomst, dan blijft de teelt open;
  - de stekbeoordeling (bakjes, wortel, plantmaat, uniformiteit, cijfer,
    opmerking) gaat mee; een beoordeling die al in de app staat blijft staan;
  - de watergift uit blad Overview gaat naar watergift_dag met bron 'excel':
    dat is de ingestelde gift, een gemeten Priva-waarde blijft altijd voor.
    Opgeteld per teelt komt die uit op de liters die in Aantekeningen staan,
    dus die kolom hoeft niet apart mee;
  - uitvalpercentage, emmers per oogstdag en het klimaat per week gaan niet
    mee: daar is in de database (nog) geen goede plek voor.

Gebruik:
    python importeer_excel_tuin1.py              # proefdraai: toont alleen
    python importeer_excel_tuin1.py --uitvoeren  # schrijft naar de database in .env

Weigert naar Supabase te schrijven tenzij --productie erbij staat. Teelten die
al bestaan (zelfde vak en startdatum) worden overgeslagen, dus nogmaals draaien
kan geen kwaad.
"""
import argparse
import io
import os
import re
import statistics
import sys
import warnings
from collections import Counter
from datetime import date, timedelta

import openpyxl

import database

HIER = os.path.dirname(os.path.abspath(__file__))
EXCEL = os.path.join(HIER, "Teelt", "Klimaatregistratie tuin 1 kopie.xlsx")
TUIN = 1
VANAF = (2025, 49)  # eerste planting van de huidige ronde op tuin 1
GEBRUIKER = "excel-import"
DAGNUMMER = {"ma": 1, "di": 2, "wo": 3, "do": 4, "vr": 5, "za": 6, "zo": 7}
STANDAARD_PLANTDAG = 1  # maandag, als de dag niet ingevuld is
OOGSTDUUR = 4  # dagen tussen eerste en laatste oogstdag, als die laatste ontbreekt
# In Excel staan de rassen anders geschreven dan de officiële naam.
RAS_NAAM = {"camaron": "Cameron", "cameron": "Cameron",
            "vegmo single": "Single vegmo", "single vegmo": "Single vegmo"}


def tekst(waarde):
    return str(waarde).strip() if waarde not in (None, "") else None


def lees_stek(wb):
    """(jaar, week, vak) -> plantdatum, aantal planten, ras en beoordeling."""
    stek = {}
    rijen = wb["Stek"].iter_rows(values_only=True)
    next(rijen)
    for rij in rijen:
        (jaar, week, dag, vak, ras, te_poten, bakjes, _uitval, _celdagen,
         wortel, plantmaat, uniformiteit, cijfer, opmerking) = (list(rij) + [None] * 14)[:14]
        if None in (jaar, week, vak) or not isinstance(jaar, (int, float)):
            continue
        if (int(jaar), int(week)) < VANAF:
            continue
        # De plantdag is de laatste maanden niet meer ingevuld; dan maandag.
        dagnr = DAGNUMMER.get(str(dag).strip().lower()[:2], STANDAARD_PLANTDAG)
        naam_ras = RAS_NAAM.get((tekst(ras) or "").lower(), tekst(ras))
        stek[(int(jaar), int(week), int(vak))] = {
            "plantdatum": date.fromisocalendar(int(jaar), int(week), dagnr),
            "planten": int(te_poten) if te_poten else None,
            "ras": naam_ras,
            "beoordeling": {
                "ras": naam_ras,
                "bakjes": float(bakjes) if isinstance(bakjes, (int, float)) else None,
                "wortel": tekst(wortel),
                "plantmaat": tekst(plantmaat),
                "uniformiteit": tekst(uniformiteit),
                "beoordeling": int(cijfer) if isinstance(cijfer, (int, float)) else None,
                "opmerking": tekst(opmerking),
            },
        }
    return stek


def lees_oogstdag(waarde, jaar, plantweek):
    """'Vrij-24' of 'do-15' -> datum. Het jaar rolt door als de week lager is."""
    gevonden = re.match(r"\s*([a-zA-Z]+)\s*-\s*(\d+)", str(waarde or ""))
    if not gevonden:
        return None
    dagnr = DAGNUMMER.get(gevonden.group(1)[:2].lower())
    week = int(gevonden.group(2))
    if dagnr is None:
        return None
    oogstjaar = jaar + 1 if week < plantweek else jaar
    try:
        return date.fromisocalendar(oogstjaar, week, dagnr)
    except ValueError:
        return None


def lees_aantekeningen(wb):
    """(jaar, plantweek) -> vakken, eerste en laatste oogstdag, lengte."""
    per_planting = {}
    rijen = wb["Aantekeningen"].iter_rows(values_only=True)
    next(rijen)
    for rij in rijen:
        rij = (list(rij) + [None] * 14)[:14]
        jaar, week, tralie = rij[0], rij[1], rij[2]
        if not isinstance(jaar, (int, float)) or not isinstance(week, (int, float)):
            continue
        jaar, week = int(jaar), int(week)
        if (jaar, week) < VANAF:
            continue
        vakken = [int(v) for v in re.findall(r"\d+", str(tralie or ""))]
        per_planting[(jaar, week)] = {
            "vakken": vakken,
            "eerste_oogst": lees_oogstdag(rij[9], jaar, week),
            "laatste_oogst": lees_oogstdag(rij[10], jaar, week),
            "lengte": float(rij[11]) if isinstance(rij[11], (int, float)) else None,
        }
    return per_planting


def lees_teeltweken(wb):
    """
    (jaar, week, vak) -> laatste jaar-week met een waarde in blad Teelt. Dat blad
    is alleen de eerste maanden bijgehouden en loopt door t/m de oogstweek; het
    vult de plantingen aan waar in Aantekeningen nog geen oogstdag staat.
    """
    rijen = wb["Teelt"].iter_rows(values_only=True)
    kop = next(rijen)
    next(rijen)
    weekkolommen = [
        (i, tuple(int(x) for x in k.split("-")))
        for i, k in enumerate(kop) if isinstance(k, str) and "-" in k
    ]
    teelten = {}
    for rij in rijen:
        week, jaar, vak, eenheid = (list(rij) + [None] * 4)[:4]
        # Elke teelt staat er twee keer in (°C en J/cm); één is genoeg.
        if eenheid != "°C" or None in (week, jaar, vak):
            continue
        gevuld = [jw for i, jw in weekkolommen if i < len(rij) and rij[i] is not None]
        if gevuld:
            teelten[(int(jaar), int(week), int(vak))] = gevuld[-1]
    if not teelten:
        return {}
    # Toen het blad niet meer werd bijgehouden hielden alle lopende teelten in
    # dezelfde week op; die liepen toen nog en zijn hier dus onbruikbaar.
    eindweken = Counter(teelten.values())
    gestopt = max(week for week, aantal in eindweken.items() if aantal >= 3)
    return {sleutel: laatst for sleutel, laatst in teelten.items() if laatst < gestopt}


def lees_watergift(wb):
    """(vak, datum) -> liter per m2 uit de "W"-rijen van blad Overview."""
    rijen = wb["Overview"].iter_rows(values_only=True)
    jaren, weken, dagen = next(rijen), next(rijen), next(rijen)
    next(rijen)
    # De kop staat in drie rijen: jaar en week alleen boven de eerste dag van
    # het blok, daarna zeven dagkolommen (het blok begint op zondag).
    kolomdatum = {}
    jaar = week = None
    positie = 0
    blokstart = None
    for i in range(6, len(jaren)):
        if jaren[i] is not None:
            jaar, week, positie = jaren[i], weken[i], 0
            try:
                blokstart = date.fromisocalendar(int(jaar), int(week), 1) - timedelta(days=1)
            except ValueError:  # bijv. week 53 in een jaar dat er 52 heeft
                blokstart = None
        elif jaar is not None:
            positie += 1
        if blokstart and i < len(dagen) and dagen[i]:
            kolomdatum[i] = blokstart + timedelta(days=positie)

    eerste_dag = date.fromisocalendar(*VANAF, 1)
    watergift = {}
    vak = None
    for rij in rijen:
        if rij[0] is not None:
            vak = rij[0]
        if rij[1] != "W" or not isinstance(vak, (int, float)):
            continue
        for i, dag in kolomdatum.items():
            waarde = rij[i] if i < len(rij) else None
            if isinstance(waarde, (int, float)) and waarde and dag >= eerste_dag:
                watergift[(int(vak), dag)] = float(waarde)
    return watergift


def bouw_teelten(stek, aantekeningen, teeltweken):
    """Eén teelt per vak, met de oogstgegevens van de planting waar het vak bij hoort."""
    # De oogst duurt een paar dagen. Waar beide dagen bekend zijn kun je zien
    # hoeveel; die mediaan gebruiken we waar de laatste dag ontbreekt.
    duren = [(g["laatste_oogst"] - g["eerste_oogst"]).days
             for g in aantekeningen.values()
             if g["eerste_oogst"] and g["laatste_oogst"]]
    oogstduur = int(statistics.median(duren)) if duren else OOGSTDUUR

    vandaag = date.today()
    teelten, geschat, zonder_oogst = [], 0, 0
    for (jaar, week, vak), gegevens in sorted(stek.items()):
        planting = aantekeningen.get((jaar, week), {})
        plantdatum = gegevens["plantdatum"]
        oogst = planting.get("laatste_oogst")
        if oogst is None and planting.get("eerste_oogst"):
            oogst = planting["eerste_oogst"] + timedelta(days=oogstduur)
            geschat += 1
        if oogst is None and (jaar, week, vak) in teeltweken:
            # Blad Teelt kent de oogst alleen per week: zelfde weekdag als geplant.
            oogstjaar, oogstweek = teeltweken[(jaar, week, vak)]
            oogst = date.fromisocalendar(oogstjaar, oogstweek, plantdatum.isoweekday())
            geschat += 1
        if oogst and oogst > vandaag:  # wordt nu geoogst: nog niet afsluiten
            oogst, geschat = None, geschat - 1
        if oogst is None:
            zonder_oogst += 1
        if planting.get("vakken") and vak not in planting["vakken"]:
            print(f"  let op: vak {vak} staat niet bij planting {jaar}-{week} "
                  f"(Aantekeningen: {planting['vakken']})")
        teelten.append({
            "vak": vak, "start": gegevens["plantdatum"], "oogst": oogst,
            "planten": gegevens["planten"], "ras": gegevens["ras"],
            "lengte": planting.get("lengte"),
        })
    return teelten, geschat, zonder_oogst, oogstduur


def bestaande_teelten(tuin_id):
    """
    De teelten die al in de database staan, als (vak, plantweek). Excel en de
    app wijken soms een dag af in de plantdatum; een vak heeft nooit twee
    teelten in één week, dus op vak + week herken je ze toch als dezelfde.
    """
    with database.get_connection() as conn:
        cur = conn.cursor()
        cur.execute("""
            SELECT v.vaknummer, t.datum_teelt_start
            FROM teelten t JOIN teeltvakken v ON v.id = t.teeltvak_id
            WHERE v.tuin_id = %s
        """, (tuin_id,))
        return {(vak, date.fromisoformat(start[:10]).isocalendar()[:2])
                for vak, start in cur.fetchall()}


def stek_bij_teelten(stek, tuin_id):
    """
    Koppelt de stekbeoordelingen uit Excel aan de teelten van deze tuin (zelfde
    vak en plantdatum) die nog geen beoordeling hebben.
    """
    with database.get_connection() as conn:
        cur = conn.cursor()
        cur.execute("""
            SELECT t.id, v.vaknummer, t.datum_teelt_start, s.id IS NOT NULL
            FROM teelten t
            JOIN teeltvakken v ON v.id = t.teeltvak_id
            LEFT JOIN stekbeoordelingen s ON s.teelt_id = t.id
            WHERE v.tuin_id = %s
        """, (tuin_id,))
        rijen = cur.fetchall()
    teelt_per_plant = {(vak, start[:10]): (tid, heeft) for tid, vak, start, heeft in rijen}
    # Excel en de app wijken soms een dag af in de plantdatum. Een vak heeft nooit
    # twee teelten in één week, dus dan mag vak + week ook.
    per_vakweek = {}
    for tid, vak, start, heeft in rijen:
        sleutel = (vak, date.fromisoformat(start[:10]).isocalendar()[:2])
        per_vakweek.setdefault(sleutel, []).append((tid, heeft))
    koppel, geen_teelt, al_ingevuld = [], 0, 0
    for (jaar, week, vak), gegevens in stek.items():
        gevonden = teelt_per_plant.get((vak, str(gegevens["plantdatum"])))
        if gevonden is None and len(per_vakweek.get((vak, (jaar, week)), [])) == 1:
            gevonden = per_vakweek[(vak, (jaar, week))][0]
        if gevonden is None:
            geen_teelt += 1
        elif gevonden[1]:
            al_ingevuld += 1
        else:
            koppel.append((gevonden[0], gegevens["beoordeling"]))
    return koppel, geen_teelt, al_ingevuld


def schrijf(teelten, tuin_id):
    with database.get_connection() as conn:
        cur = conn.cursor()
        for t in teelten:
            ras = t["ras"] if t["ras"] and t["ras"] != database.STANDAARD_RAS else None
            cur.execute("""
                INSERT INTO teelten (teeltvak_id, datum_teelt_start, datum_oogst,
                                     aantal_planten, code, ras, lengte_eind)
                VALUES (%s, %s, %s, %s, %s, %s, %s)
            """, (
                database.get_of_maak_teeltvak(t["vak"], tuin_id=tuin_id),
                str(t["start"]), str(t["oogst"]) if t["oogst"] else None,
                t["planten"], database.genereer_teelt_code(t["start"], t["vak"]),
                ras, t["lengte"],
            ))
        conn.commit()


def main():
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("--bestand", default=EXCEL, help="ander pad naar het Excel-bestand")
    parser.add_argument("--uitvoeren", action="store_true", help="echt wegschrijven (anders proefdraai)")
    parser.add_argument("--productie", action="store_true", help="toestaan dat er naar Supabase geschreven wordt")
    args = parser.parse_args()

    doel = os.environ.get("DATABASE_URL", "")
    if args.uitvoeren and "supabase" in doel and not args.productie:
        sys.exit("GESTOPT: .env wijst naar Supabase (productie). Voeg --productie toe als dat de bedoeling is.")

    tuin_id = database.get_tuin_id(TUIN)
    if tuin_id is None:
        sys.exit(f"GESTOPT: tuin {TUIN} bestaat nog niet in de database.")
    # Alles in dit script gaat over tuin 1, ook de functies zonder expliciete tuin.
    database.zet_actieve_tuin(tuin_id)

    warnings.filterwarnings("ignore", module="openpyxl")
    # Via het geheugen, want OneDrive geeft het bestand niet altijd vrij om te
    # openen terwijl kopiëren wel mag (bijv. als het in Excel openstaat).
    with open(args.bestand, "rb") as bestand:
        wb = openpyxl.load_workbook(io.BytesIO(bestand.read()), read_only=True, data_only=True)
    stek = lees_stek(wb)
    aantekeningen = lees_aantekeningen(wb)
    teeltweken = lees_teeltweken(wb)
    teelten, geschat, zonder_oogst, oogstduur = bouw_teelten(stek, aantekeningen, teeltweken)

    bestaand = bestaande_teelten(tuin_id)
    nieuw = [t for t in teelten
             if (t["vak"], t["start"].isocalendar()[:2]) not in bestaand]

    print(f"Plantingen in Excel: {len(aantekeningen)}, teelten (vak-plantingen): {len(teelten)}, "
          f"nieuw: {len(nieuw)}, stond al in de database: {len(teelten) - len(nieuw)}")
    print(f"Oogstdatum: {len(teelten) - zonder_oogst} bekend "
          f"(waarvan {geschat} geschat uit de eerste oogstdag of de oogstweek), "
          f"{zonder_oogst} nog lopend")
    print("Rassen: " + ", ".join(f"{r} {a}x" for r, a in Counter(t["ras"] for t in teelten).items()))
    if nieuw:
        duren = [(t["oogst"] - t["start"]).days for t in nieuw if t["oogst"]]
        print(f"Periode: gestart {min(t['start'] for t in nieuw):%d-%m-%y} t/m "
              f"{max(t['start'] for t in nieuw):%d-%m-%y}"
              + (f", teeltduur {min(duren)}-{max(duren)} dagen" if duren else ""))
        for t in nieuw[:3] + nieuw[-3:]:
            oogst = f"{t['oogst']:%d-%m-%y}" if t["oogst"] else "loopt nog"
            print(f"  vak {t['vak']:>2}  start {t['start']:%d-%m-%y}  oogst {oogst:>9}  "
                  f"{t['planten']} planten  {t['ras']}")
    lopend = [t for t in teelten if not t["oogst"]]
    if lopend:
        print("Nog open na de import (afsluiten in de app als ze toch geoogst zijn):")
        for t in sorted(lopend, key=lambda t: (t["start"], t["vak"])):
            print(f"  vak {t['vak']:>2}  gestart {t['start']:%d-%m-%y}")

    if args.uitvoeren and nieuw:
        schrijf(nieuw, tuin_id)
        database.log_wijziging(
            GEBRUIKER, "aangemaakt", "teelt", None,
            f"{len(nieuw)} teelten van tuin 1 uit Excel geïmporteerd "
            f"(gestart {min(t['start'] for t in nieuw):%d-%m-%y} t/m "
            f"{max(t['start'] for t in nieuw):%d-%m-%y})"
        )
        print(f"\n{len(nieuw)} teelten geïmporteerd.")

    # Pas na de teelten: een beoordeling hangt aan een teelt die nu pas bestaat.
    koppel, geen_teelt, al_ingevuld = stek_bij_teelten(stek, tuin_id)
    print(f"\nStekbeoordelingen: {len(koppel)} nieuw, {al_ingevuld} stonden er al, "
          f"{geen_teelt} zonder bijbehorende teelt in de database")

    watergift = lees_watergift(wb)
    if watergift:
        dagen_wg = sorted({dag for _, dag in watergift})
        vakken_wg = sorted({vak for vak, _ in watergift})
        print(f"Watergift (ingesteld): {len(watergift)} vak-dagen, vak {min(vakken_wg)}-{max(vakken_wg)}, "
              f"{dagen_wg[0]:%d-%m-%y} t/m {dagen_wg[-1]:%d-%m-%y}")

    if not args.uitvoeren:
        print("\nProefdraai: er is niets geschreven. Gebruik --uitvoeren om te importeren.")
        return

    for teelt_id, velden in koppel:
        database.sla_stekbeoordeling_op(teelt_id, velden, gebruiker=GEBRUIKER)
    if koppel:
        print(f"{len(koppel)} stekbeoordelingen geïmporteerd.")

    for (vak, dag), liter in sorted(watergift.items()):
        database.upsert_watergift_dag(vak, dag, liter, bron="excel", tuin_id=tuin_id)
    print(f"{len(watergift)} vak-dagen watergift weggeschreven (bron excel; "
          "gemeten Priva-waarden blijven staan).")


if __name__ == "__main__":
    main()
